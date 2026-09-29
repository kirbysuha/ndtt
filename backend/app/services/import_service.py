"""
Excel import servisi.
Mevcut 'Modül-1 Not Durum Takip Tablosu.xlsx' dosyasını sisteme aktarır.
"""
import io
from datetime import datetime, date
import openpyxl
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from app.models.lesson import Lesson
from app.models.note_tracking import NoteTracking, NoteStatus
from app.models.status_history import StatusHistory
from app.services.timer_service import now_local


# Excel'deki durum metinleri → NoteStatus enum eşlemesi
STATUS_MAP = {
    "Not henüz ulaşmadı": NoteStatus.NOT_REACHED,
    "Not mailde": NoteStatus.IN_MAIL,
    "Not denetimde": NoteStatus.REVIEW,
    "Not 2. Kez denetimde": NoteStatus.REVIEW_2,
    "Not kabul edildi": NoteStatus.ACCEPTED,
    "Not 1. Kez reddedildi": NoteStatus.REJECTED_1,
    "Not 2. Kez reddedildi": NoteStatus.REJECTED_2,
    "Not yüklendi": NoteStatus.UPLOADED,
}


def parse_excel_file(file_content: bytes) -> list[dict]:
    """
    Excel dosyasını parse et ve satır listesi döndür.
    Sütun sırası (Apps Script ile aynı):
      Col 1  = Son Güncelleme
      Col 2  = Dersin Adı
      Col 3  = Dersin Sırası (s1, s2...)
      Col 4  = Dersin Konusu
      Col 5  = Ders Tarihi
      Col 6  = Hazırlayan Ekip
      Col 7  = Notun Durumu
      Col 10 = 1. Reddedilme Tarihi (önemli, saklanır)
      Col 12 = 2. Reddedilme Tarihi (önemli, saklanır)
      Col 14 = Notun Yüklenme Süresi (dondurulan süre, saklanır)
      Col 15 = Dersin Ses Kaydının Durumu
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
    ws = wb.active

    rows = []
    for row_idx in range(2, ws.max_row + 1):
        row = ws[row_idx]

        def cell(col: int):
            return row[col - 1].value

        subject_name = cell(2)
        if not subject_name:
            continue

        status_text = cell(7) or "Not henüz ulaşmadı"
        status = STATUS_MAP.get(str(status_text).strip(), NoteStatus.NOT_REACHED)

        lesson_date_raw = cell(5)
        lesson_date = None
        if isinstance(lesson_date_raw, datetime):
            lesson_date = lesson_date_raw.date()
        elif isinstance(lesson_date_raw, date):
            lesson_date = lesson_date_raw

        first_reject_raw = cell(10)
        first_reject_date = first_reject_raw if isinstance(first_reject_raw, datetime) else None

        second_reject_raw = cell(12)
        second_reject_date = second_reject_raw if isinstance(second_reject_raw, datetime) else None

        upload_duration = cell(14)
        if upload_duration and str(upload_duration).strip() in ("-", ""):
            upload_duration = None

        rows.append({
            "subject_name": str(subject_name).strip(),
            "order_label": str(cell(3)).strip() if cell(3) else None,
            "topic": str(cell(4)).strip() if cell(4) else None,
            "lesson_date": lesson_date,
            "assigned_team": str(cell(6)).strip() if cell(6) else None,
            "status": status,
            "first_reject_date": first_reject_date,
            "second_reject_date": second_reject_date,
            "final_upload_duration": str(upload_duration).strip() if upload_duration else None,
            "audio_status": str(cell(15)).strip() if cell(15) else None,
        })

    return rows


async def import_excel_to_module(
    db: AsyncSession,
    module_id: str,
    file_content: bytes,
    overwrite: bool = False,
) -> dict:
    """Excel dosyasını parse edip belirtilen modüle import et."""
    rows = parse_excel_file(file_content)
    now = now_local()

    created_count = 0
    updated_count = 0
    skipped_count = 0

    for row_data in rows:
        result = await db.execute(
            select(Lesson).where(
                Lesson.module_id == module_id,
                Lesson.subject_name == row_data["subject_name"],
                Lesson.order_label == row_data["order_label"],
            )
        )
        lesson = result.scalar_one_or_none()

        if lesson and not overwrite:
            skipped_count += 1
            continue

        if not lesson:
            lesson = Lesson(
                module_id=module_id,
                subject_name=row_data["subject_name"],
                order_label=row_data["order_label"],
                topic=row_data["topic"],
                lesson_date=row_data["lesson_date"],
                assigned_team=row_data["assigned_team"],
                audio_status=row_data["audio_status"],
            )
            db.add(lesson)
            await db.flush()
            created_count += 1
        else:
            lesson.topic = row_data["topic"]
            lesson.lesson_date = row_data["lesson_date"]
            lesson.assigned_team = row_data["assigned_team"]
            lesson.audio_status = row_data["audio_status"]
            updated_count += 1

        tracking_result = await db.execute(
            select(NoteTracking).where(NoteTracking.lesson_id == lesson.id)
        )
        tracking = tracking_result.scalar_one_or_none()

        if not tracking:
            tracking = NoteTracking(lesson_id=lesson.id)
            db.add(tracking)
            await db.flush()

        old_status = tracking.status
        tracking.status = row_data["status"]
        tracking.first_reject_date = row_data["first_reject_date"]
        tracking.second_reject_date = row_data["second_reject_date"]
        tracking.final_upload_duration = row_data["final_upload_duration"]

        if old_status != row_data["status"]:
            history = StatusHistory(
                note_tracking_id=tracking.id,
                old_status=old_status,
                new_status=row_data["status"],
                changed_by="excel_import",
                changed_at=now,
            )
            db.add(history)

    return {
        "total": len(rows),
        "created": created_count,
        "updated": updated_count,
        "skipped": skipped_count,
    }
